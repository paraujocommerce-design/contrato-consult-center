import streamlit as st
from docx import Document
import io
import re
from datetime import datetime

# --- CONFIGURAÇÃO E LIMPEZA ---
st.set_page_config(page_title="Consult Center - Emissor", layout="wide")

def limpar_campos():
    for key in st.session_state.keys():
        del st.session_state[key]
    st.rerun()

st.title("📄 Emissor de Contrato Profissional")
st.button("🆕 Novo Cliente (Limpar Tudo)", on_click=limpar_campos)

# Inicialização de variáveis de controle
if 'arquivo_gerado' not in st.session_state:
    st.session_state.arquivo_gerado = None

# --- FORMULÁRIO ---
with st.form("form_v3"):
    c1, c2, c3 = st.columns(3)
    with c1:
        st.subheader("🏢 Identificação")
        razao = st.text_input("Razão Social")
        fantasia = st.text_input("Nome Fantasia")
        cnpj = st.text_input("CNPJ (números)")
        rep = st.text_input("Representante")
    with c2:
        st.subheader("📍 Localização")
        end = st.text_input("Endereço")
        bairro = st.text_input("Bairro")
        cidade = st.text_input("Cidade")
        uf = st.text_input("UF")
        cep = st.text_input("CEP")
        ref = st.text_input("Ponto de Referência")
    with c3:
        st.subheader("📞 Contatos")
        email = st.text_input("E-mail")
        ddd = st.text_input("DDD")
        tel = st.text_input("Telefone")
        cel1 = st.text_input("Celular 01")
        cel2 = st.text_input("Celular 02")

    st.divider()
    st.subheader("📊 Quantitativos e Financeiro")
    v1, v2, v3, v4 = st.columns(4)
    with v1: q01 = st.number_input("Qtd Serasa (Op 01)", min_value=0)
    with v2: q13 = st.number_input("Qtd Negativação (Op 13)", min_value=0)
    with v3: q05 = st.number_input("Qtd Localizador (Op 05)", min_value=0)
    with v4: valor = st.text_input("Valor Mensal R$")

    st.subheader("➕ Outros Serviços")
    o1, o2 = st.columns(2)
    with o1: op_nome = st.text_input("Número da Opção")
    with o2: qout = st.number_input("Quantidade Outros", min_value=0)

    st.divider()
    st.subheader("🛡️ Garantia")
    f1, f2, f3 = st.columns(3)
    with f1: fiador = st.text_input("Fiador")
    with f2: cpf_f = st.text_input("CPF Fiador")
    with f3: resp_p = st.text_input("Responsável Pagamento")

    submit = st.form_submit_button("VALIDAR E GERAR CONTRATO")

if submit:
    try:
        doc = Document("CONTRATO.docx")
        data_hoje = datetime.now().strftime("%d/%m/%Y")
        
        # Mapa usando etiquetas curtas para proteger a diagramação
        trocas = {
            "{{RAZAO}}": razao, "{{FANTASIA}}": fantasia, "{{CNPJ}}": cnpj,
            "{{REP}}": rep, "{{ENDERECO}}": end, "{{BAIRRO}}": bairro,
            "{{CIDADE}}": cidade, "{{UF}}": uf, "{{CEP}}": cep, "{{REF}}": ref,
            "{{EMAIL}}": email, "{{D1}}": ddd, "{{D2}}": ddd, "{{D3}}": ddd,
            "{{TEL}}": tel, "{{CEL1}}": cel1, "{{CEL2}}": cel2,
            "{{VALOR}}": valor, "{{FIADOR}}": fiador, "{{CPF}}": cpf_f,
            "{{RP}}": resp_p, "{{Q01}}": str(q01), "{{Q13}}": str(q13),
            "{{Q05}}": str(q05), "{{OP}}": op_nome, "{{QOUT}}": str(qout),
            "{{DATA}}": data_hoje
        }

        for tabela in doc.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    for paragrafo in celula.paragraphs:
                        for tag, info in trocas.items():
                            if tag in paragrafo.text:
                                paragrafo.text = paragrafo.text.replace(tag, info)

        output = io.BytesIO()
        doc.save(output)
        st.session_state.arquivo_gerado = output.getvalue()
        st.session_state.nome_doc = razao
        st.success(f"✅ Contrato processado para {razao}!")
    except Exception as e:
        st.error(f"Erro: {e}")

if st.session_state.arquivo_gerado:
    st.download_button(
        label="📥 BAIXAR CONTRATO AGORA",
        data=st.session_state.arquivo_gerado,
        file_name=f"Contrato_{st.session_state.nome_doc}.docx"
    )
